#!/usr/bin/env python3
"""8-slide deck: how many bones/fractures this project can test."""

from __future__ import annotations

import json
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent
ATLAS = json.loads((ROOT / "bone_coverage.json").read_text())
OUT = ROOT / "Bone_Coverage_Presentation.pptx"

NAVY = RGBColor(7, 20, 28)
TEAL = RGBColor(46, 196, 182)
WHITE = RGBColor(244, 251, 252)
MUTED = RGBColor(176, 204, 214)
CARD = RGBColor(16, 42, 56)
W, H = Inches(13.333), Inches(7.5)


def run(p, text, size=16, bold=False, color=WHITE):
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    r.font.name = "Calibri"


def box(slide, l, t, w, h, fill=CARD):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.fill.background()
    try:
        sh.adjustments[0] = 0.08
    except Exception:
        pass
    return sh


def bg(slide):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H)
    sh.fill.solid()
    sh.fill.fore_color.rgb = NAVY
    sh.line.fill.background()
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, Inches(0.1))
    bar.fill.solid()
    bar.fill.fore_color.rgb = TEAL
    bar.line.fill.background()


def heading(slide, kicker, title):
    a = slide.shapes.add_textbox(Inches(0.45), Inches(0.22), Inches(12.4), Inches(0.28))
    tf = a.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    run(p, kicker.upper(), 12, True, TEAL)
    b = slide.shapes.add_textbox(Inches(0.45), Inches(0.48), Inches(12.4), Inches(0.5))
    tf = b.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    run(p, title, 26, True, WHITE)


def put(slide, l, t, w, h, lines, size=14):
    tb = slide.shapes.add_textbox(l, t, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.clear()
    first = True
    for line, *rest in ((ln,) if isinstance(ln, str) else ln for ln in lines):
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        bold = rest[0] if rest else False
        run(p, line, size, bold, WHITE if not first or True else WHITE)
        p.space_after = Pt(6)
    return tf


def build():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    blank = prs.slide_layouts[6]

    s = prs.slides.add_slide(blank)
    bg(s)
    heading(s, "Coverage  ·  Viva slide", "How many fractures of the human body can we test?")
    metrics = [
        ("206", "Bones in the adult skeleton (we do not classify all of them)"),
        ("10", "Morphological fracture types (IEEE Access 2025)"),
        ("11", "Skeletal regions in this project"),
        ("~42", "Named bones / bone groups inside those regions"),
        ("7", "Upper-limb YOLO classes in the published 86% mAP50 split"),
        ("110", "Max region × morphology combinations (not all labelled yet)"),
    ]
    for i, (n, cap) in enumerate(metrics):
        col, row = i % 3, i // 3
        l, t = Inches(0.45 + col * 4.2), Inches(1.25 + row * 2.7)
        box(s, l, t, Inches(4.0), Inches(2.45))
        tb = s.shapes.add_textbox(l, t + Inches(0.35), Inches(4.0), Inches(0.7))
        tf = tb.text_frame
        tf.clear()
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run(p, n, 32, True, TEAL)
        capb = s.shapes.add_textbox(l + Inches(0.2), t + Inches(1.15), Inches(3.6), Inches(1.05))
        tf = capb.text_frame
        tf.word_wrap = True
        tf.clear()
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run(p, cap, 13, False, MUTED)

    s = prs.slides.add_slide(blank)
    bg(s)
    heading(s, "Do not mix these two questions", "Pattern of break  vs  bone / region")
    box(s, Inches(0.45), Inches(1.25), Inches(6.1), Inches(5.5))
    put(
        s,
        Inches(0.65),
        Inches(1.4),
        Inches(5.7),
        Inches(5.2),
        [
            ("1. Morphology — HOW it broke", True),
            "Avulsion, comminuted, dislocation, greenstick, hairline, impacted, longitudinal, oblique, pathological, spiral.",
            "IEEE 10-class set: 1,129 X-rays. VGG-16 / VGG-16+RF accuracy 95%.",
            "",
            ("2. Anatomy — WHERE it is", True),
            "Shoulder, humerus, elbow, forearm, wrist, fingers, pelvis, femur, knee, ankle, spine.",
            "Abstract YOLO: 7 upper-limb labels, 3,316 / 399 images, mAP50 86%.",
        ],
        15,
    )
    box(s, Inches(6.8), Inches(1.25), Inches(6.05), Inches(5.5))
    put(
        s,
        Inches(7.0),
        Inches(1.4),
        Inches(5.7),
        Inches(5.2),
        [
            ("Spoken answer", True),
            ATLAS["headline_for_viva"],
            "",
            "A complete report line is: “spiral fracture of the tibial shaft” = type + bone. The app outputs type and region; it does not name every carpal or every vertebra separately.",
        ],
        15,
    )

    s = prs.slides.add_slide(blank)
    bg(s)
    heading(s, "Upper limb  ·  published detector", "Bones in the 86% mAP50 YOLO set")
    rows = [
        "Shoulder — clavicle, scapula, proximal humerus",
        "Humerus — shaft and distal humerus",
        "Elbow — distal humerus, radial head/neck, olecranon, coronoid",
        "Forearm — radius and ulna shafts",
        "Wrist — distal radius/ulna, scaphoid and other carpals",
        "Fingers — metacarpals 1–5 and all phalanges",
    ]
    for i, line in enumerate(rows):
        t = Inches(1.22 + i * 0.9)
        box(s, Inches(0.5), t, Inches(12.3), Inches(0.78))
        put(s, Inches(0.7), t + Inches(0.18), Inches(12.0), Inches(0.5), [line], 18)

    s = prs.slides.add_slide(blank)
    bg(s)
    heading(s, "Whole-body scale-up", "Extra regions in software (need more labelled films)")
    extra = [
        ("Pelvis", "Ilium, ischium, pubis, sacrum, acetabulum"),
        ("Femur", "Neck, intertrochanteric region, shaft, distal femur"),
        ("Knee", "Distal femur, patella, tibial plateau, proximal fibula"),
        ("Ankle", "Distal tibia & fibula (malleoli), talus"),
        ("Spine", "Cervical, thoracic, lumbar vertebrae as groups — not 24 IDs"),
        ("20k story", "GRAZPEDWRI-DX = 20,327 wrist studies (not 20k of every bone)"),
    ]
    for i, (title, body) in enumerate(extra):
        col, row = i % 3, i // 3
        l, t = Inches(0.45 + col * 4.2), Inches(1.25 + row * 2.85)
        box(s, l, t, Inches(4.0), Inches(2.65))
        put(s, Inches(l + Inches(0.18)), t + Inches(0.2), Inches(3.65), Inches(2.3), [(title, True), body], 15)

    s = prs.slides.add_slide(blank)
    bg(s)
    heading(s, "Ten patterns", "Morphology classes (any region)")
    morph = ATLAS["morphology"]
    for i, m in enumerate(morph):
        col, row = i % 5, i // 5
        l, t = Inches(0.35 + col * 2.58), Inches(1.2 + row * 2.85)
        box(s, l, t, Inches(2.45), Inches(2.65))
        put(s, l + Inches(0.12), t + Inches(0.15), Inches(2.22), Inches(2.35), [(m["name"].replace(" Fracture", ""), True), m["plain"]], 12)

    s = prs.slides.add_slide(blank)
    bg(s)
    heading(s, "Out of scope as dedicated classes", "What we do not claim")
    put(
        s,
        Inches(0.6),
        Inches(1.3),
        Inches(12.1),
        Inches(5.5),
        [
            "Skull and facial bones",
            "Ribs and sternum",
            "Most bones of the foot (metatarsals / midfoot) as their own classes",
            "Each of the 8 carpals or 24 vertebrae as separate outputs (they are grouped)",
            "Replacing a radiologist, or using the metal overlay as a surgical plan",
            "",
            "Those body parts can be added later with new labels and images — they are not in the current class list.",
        ],
        18,
    )

    s = prs.slides.add_slide(blank)
    bg(s)
    heading(s, "Data you can show", "Counts behind the demo")
    lines = [
        "10-class Kaggle morphology set — 1,129 images (paper Table 1).",
        "YOLOv8 abstract split — 3,316 train / 399 val, 7 upper-limb classes.",
        "GRAZPEDWRI-DX — 20,327 pediatric wrist rows (CSV in the repo, match verified).",
        "FracAtlas — 4,083 mixed musculoskeletal X-rays (YOLO/COCO/VOC).",
        "Live check: python main.py audit --probe  and  /datasets in the web app.",
        "Bone map in the app: /coverage  (Body coverage).",
    ]
    for i, line in enumerate(lines):
        t = Inches(1.25 + i * 0.85)
        box(s, Inches(0.5), t, Inches(12.3), Inches(0.72))
        put(s, Inches(0.7), t + Inches(0.16), Inches(12.0), Inches(0.48), [line], 16)

    s = prs.slides.add_slide(blank)
    bg(s)
    heading(s, "Close this mini-deck", "One sentence to remember")
    box(s, Inches(0.7), Inches(2.0), Inches(11.9), Inches(3.6))
    put(
        s,
        Inches(1.0),
        Inches(2.3),
        Inches(11.3),
        Inches(3.1),
        [
            "We test 10 ways a bone can break, on X-rays from 11 body regions — about 42 named bones — not all 206 bones of the human body.",
            "",
            "Published scores: 95% morphology (VGG-16 family), 86% mAP50 upper-limb detection.",
            "",
            "SHIBIN ANTONY  ·  811241013",
        ],
        18,
    )

    prs.save(OUT)
    print("Wrote", OUT)


if __name__ == "__main__":
    build()
